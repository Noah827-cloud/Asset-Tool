"""Utilities for writing extracted items to Excel workbooks."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import EquipmentItem


DEFAULT_HEADERS = [
    "设备名称",
    "规格/型号",
    "数量",
    "单位",
    "单价",
    "合计",
    "来源文件",
    "页码",
    "原始行数据",
]


class ExcelWriter:
    """Persist equipment items into an Excel template."""

    def __init__(self, template: Optional[Path | str] = None, sheet_name: str = "设备清单") -> None:
        self.template = Path(template) if template else None
        self.sheet_name = sheet_name

    def write(self, items: Iterable[EquipmentItem], output_path: Path | str) -> None:
        workbook, worksheet = self._prepare_workbook()
        for row in self._ensure_headers(worksheet):
            worksheet.append(row)
        for item in items:
            worksheet.append(item.as_row())
        workbook.save(output_path)

    def _prepare_workbook(self) -> tuple[Workbook, Worksheet]:
        if self.template and self.template.exists():
            workbook = load_workbook(self.template)
        else:
            workbook = Workbook()
        worksheet = workbook[self.sheet_name] if self.sheet_name in workbook.sheetnames else workbook.active
        worksheet.title = self.sheet_name
        return workbook, worksheet

    def _ensure_headers(self, worksheet: Worksheet) -> List[List[str]]:
        if worksheet.max_row == 1 and all(cell.value is None for cell in worksheet[1]):
            return [DEFAULT_HEADERS]
        first_row_values = [cell.value for cell in worksheet[1]]
        if not first_row_values or any(value is None for value in first_row_values):
            return [DEFAULT_HEADERS]
        return []


__all__ = ["ExcelWriter", "DEFAULT_HEADERS"]
